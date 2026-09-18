"""
assembler.py — Ghép storyline + ảnh → HTML slides + PPTX
"""
import os
import json
import base64
from pathlib import Path
from jinja2 import Environment, BaseLoader


# ─── Brand defaults (Ahamove) — override qua config ──────────────────────────
BRAND = {
    "primary": "#0E4174",
    "accent": "#FF7F32",
    "success": "#10B981",
    "danger": "#EF4444",
    "bg": "#F8FAFC",
    "white": "#FFFFFF",
    "text_dark": "#0F172A",
    "text_muted": "#64748B",
    "font": "Lexend",
}

SLIDE_W = 1920
SLIDE_H = 1080


# ─── Main entry ───────────────────────────────────────────────────────────────

def assemble_html(analysis: dict, images_dir: str, output_path: str, brand: dict = None):
    """Tạo HTML file chứa tất cả slides (1920×1080 mỗi slide)."""
    b = brand or BRAND
    slides_html = []

    for slide in analysis["output_slides"]:
        img_path = _get_image_path(slide, images_dir)
        slide_html = _render_slide(slide, img_path, b)
        slides_html.append(slide_html)

    html = _wrap_html(slides_html, analysis, b)

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html)

    print(f"[assemble] HTML saved: {output_path}")
    return output_path


def assemble_pptx(analysis: dict, images_dir: str, output_path: str, brand: dict = None):
    """Tạo PPTX file từ analysis JSON + images."""
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    b = brand or BRAND
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W * 9525)   # EMU = px * 9525
    prs.slide_height = Emu(SLIDE_H * 9525)

    blank_layout = prs.slide_layouts[6]  # blank layout

    for slide_data in analysis["output_slides"]:
        slide = prs.slides.add_slide(blank_layout)
        _add_pptx_slide(slide, slide_data, images_dir, b, prs)

    prs.save(output_path)
    print(f"[assemble] PPTX saved: {output_path}")
    return output_path


# ─── HTML renderers ───────────────────────────────────────────────────────────

def _render_slide(slide: dict, img_path: str | None, b: dict) -> str:
    layout = slide.get("layout_type", "executive_summary")
    renderer = LAYOUT_RENDERERS.get(layout, _render_default)
    return renderer(slide, img_path, b)


def _wrap_html(slides_html: list[str], analysis: dict, b: dict) -> str:
    slides_joined = "\n".join(
        f'<div class="slide" id="slide-{i+1}">{s}</div>'
        for i, s in enumerate(slides_html)
    )
    title = analysis.get("report_title", "Report")
    total_in = analysis.get("total_input_slides", "?")
    total_out = len(analysis.get("output_slides", []))

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Lexend:wght@300;400;500;600;700;800;900&display=swap" rel="stylesheet">
<style>
  *, *::before, *::after {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{
    font-family: '{b["font"]}', system-ui, sans-serif;
    background: #111827;
    color: {b["text_dark"]};
  }}
  .controls {{
    position: fixed; top: 16px; left: 50%; transform: translateX(-50%);
    z-index: 100; display: flex; gap: 12px; align-items: center;
    background: rgba(0,0,0,.7); padding: 10px 20px; border-radius: 40px;
    backdrop-filter: blur(8px);
  }}
  .controls button {{
    background: {b["primary"]}; color: white; border: none;
    padding: 8px 20px; border-radius: 20px; cursor: pointer; font-size: 14px;
    font-family: '{b["font"]}', sans-serif;
  }}
  .controls button:hover {{ background: {b["accent"]}; }}
  .slide-counter {{
    color: white; font-size: 14px; min-width: 80px; text-align: center;
  }}
  .slide-nav {{ display: flex; gap: 8px; flex-wrap: wrap; justify-content: center;
    padding: 0 20px; }}
  .slide-nav span {{
    cursor: pointer; color: #9CA3AF; font-size: 12px;
    padding: 4px 10px; border-radius: 12px; background: rgba(255,255,255,.1);
  }}
  .slide-nav span.active, .slide-nav span:hover {{
    background: {b["primary"]}; color: white;
  }}
  .slides-container {{ padding-top: 72px; }}
  .slide {{
    width: {SLIDE_W}px; height: {SLIDE_H}px;
    margin: 24px auto; display: none; position: relative; overflow: hidden;
    border-radius: 12px; box-shadow: 0 20px 60px rgba(0,0,0,.5);
  }}
  .slide.active {{ display: flex; }}
  /* Print / export: show all */
  @media print {{
    .controls {{ display: none; }}
    .slide {{ display: block !important; page-break-after: always;
      margin: 0; border-radius: 0; box-shadow: none; }}
    body {{ background: white; }}
  }}
</style>
</head>
<body>
<div class="controls">
  <button onclick="prevSlide()">← Prev</button>
  <span class="slide-counter" id="counter">1 / {total_out}</span>
  <button onclick="nextSlide()">Next →</button>
  <button onclick="window.print()">Export PDF</button>
  <span style="color:#9CA3AF;font-size:12px">{total_in} slides → {total_out} slides</span>
</div>
<div class="slides-container" id="container">
  {slides_joined}
</div>
<script>
  let current = 0;
  const slides = document.querySelectorAll('.slide');
  const counter = document.getElementById('counter');
  function showSlide(n) {{
    slides.forEach(s => s.classList.remove('active'));
    slides[n].classList.add('active');
    counter.textContent = (n+1) + ' / ' + slides.length;
  }}
  function nextSlide() {{ if (current < slides.length-1) showSlide(++current); }}
  function prevSlide() {{ if (current > 0) showSlide(--current); }}
  document.addEventListener('keydown', e => {{
    if (e.key === 'ArrowRight' || e.key === 'ArrowDown') nextSlide();
    if (e.key === 'ArrowLeft' || e.key === 'ArrowUp') prevSlide();
  }});
  showSlide(0);
</script>
</body>
</html>"""


# ─── Layout renderers ─────────────────────────────────────────────────────────

def _render_kpi_cards(slide: dict, img: str | None, b: dict) -> str:
    content = slide.get("content", {})

    # Flatten bất kỳ structure nào thành list of {label, value, delta, note}
    kpis = _extract_kpis(content)
    table_html = _extract_table(content, b)

    kpi_html = ""
    for kpi in kpis[:6]:
        delta = str(kpi.get("delta", ""))
        delta_color = b["success"] if delta.startswith("+") else (b["danger"] if delta.startswith("-") else "rgba(255,255,255,.5)")
        kpi_html += f"""
        <div style="background:rgba(255,255,255,.1);border-radius:20px;padding:24px 22px;flex:1;min-width:160px;backdrop-filter:blur(4px)">
          <div style="color:rgba(255,255,255,.6);font-size:11px;font-weight:600;margin-bottom:10px;letter-spacing:.06em;text-transform:uppercase">{kpi.get('label','')}</div>
          <div style="font-size:44px;font-weight:900;color:white;line-height:1;letter-spacing:-.02em">{kpi.get('value','—')}</div>
          {f'<div style="margin-top:8px;font-size:14px;font-weight:700;color:{delta_color}">{delta}</div>' if delta else ""}
          {f'<div style="margin-top:4px;font-size:12px;color:rgba(255,255,255,.5)">{kpi["note"]}</div>' if kpi.get("note") else ""}
        </div>"""

    return f"""
    <div style="width:100%;height:100%;background:linear-gradient(135deg,{b['primary']} 0%,#0a2a52 100%);
                display:flex;flex-direction:column;padding:48px 64px">
      <div style="font-size:11px;font-weight:700;color:{b['accent']};letter-spacing:.15em;text-transform:uppercase;margin-bottom:12px">KPI</div>
      <div style="font-size:30px;font-weight:800;color:white;line-height:1.25;margin-bottom:8px;max-width:900px">
        {slide.get('slide_title','')}
      </div>
      <div style="font-size:13px;color:rgba(255,255,255,.55);margin-bottom:28px">{slide.get('key_message','')}</div>
      {f'<div style="display:flex;gap:14px;flex-wrap:wrap;margin-bottom:{"24px" if table_html else "0"}">{kpi_html}</div>' if kpi_html else ""}
      {table_html}
    </div>"""


def _extract_kpis(content: dict) -> list[dict]:
    """Flatten nested content → list of simple KPI dicts."""
    # Trường hợp 1: flat kpis list
    if "kpis" in content and isinstance(content["kpis"], list):
        return content["kpis"]

    # Trường hợp 2: kpi_summary list (Retention)
    if "kpi_summary" in content:
        result = []
        for m in content["kpi_summary"]:
            market = m.get("market", "")
            rr = m.get("rr_overall", m.get("value", ""))
            mom = m.get("mom", "")
            result.append({"label": f"RR {market}", "value": rr, "delta": mom})
        return result

    # Trường hợp 3: supply_hours → lấy gap_vs_plan của từng market
    if "supply_hours" in content:
        result = []
        for market, data in content["supply_hours"].items():
            if isinstance(data, dict):
                result.append({
                    "label": f"Supply Hours {market}",
                    "value": f'{data.get("total_sh",""):,}',
                    "delta": data.get("gap_vs_plan", ""),
                    "note": "vs plan"
                })
                # Thêm demand metrics
                dm = data.get("demand_metrics", {})
                for metric, vals in dm.items():
                    if isinstance(vals, dict):
                        result.append({
                            "label": f"{metric} {market}",
                            "value": vals.get("value", ""),
                            "delta": vals.get("wow", ""),
                            "note": "WoW"
                        })
        return result[:6]

    # Trường hợp 4: amnesty / capacity
    if "amnesty_program" in content:
        a = content["amnesty_program"]
        result = [
            {"label": "Ân Xá Đạt", "value": f'{a.get("achieved","")}', "delta": f'Target {a.get("target","")}', "note": f'{a.get("achievement_rate","")}'},
        ]
        cap = content.get("capacity_han_m8", {})
        if cap:
            result.append({"label": "Capacity HAN M8", "value": cap.get("overall_achievement",""), "delta": cap.get("mtd_gap",""), "note": cap.get("root_cause","")[:40]})
        return result

    # Fallback: dùng tất cả key-value string đầu tiên
    result = []
    for k, v in content.items():
        if isinstance(v, (str, int, float)):
            result.append({"label": k, "value": str(v), "delta": "", "note": ""})
    return result[:6]


def _extract_table(content: dict, b: dict) -> str:
    """Tạo table HTML từ segment_breakdown hoặc cancel_rate."""
    # Retention segment breakdown
    if "segment_breakdown" in content:
        sb = content["segment_breakdown"]
        rows = ""
        for market, segs in sb.items():
            for seg in segs:
                mom = str(seg.get("mom", ""))
                mom_color = b["success"] if mom.startswith("+") else b["danger"]
                rows += f"""<tr>
                  <td style="padding:10px 14px;font-size:13px;color:rgba(255,255,255,.5)">{market}</td>
                  <td style="padding:10px 14px;font-size:13px;font-weight:700;color:white">{seg.get('segment','')}</td>
                  <td style="padding:10px 14px;font-size:13px;color:white">{seg.get('rr','')}</td>
                  <td style="padding:10px 14px;font-size:13px;font-weight:700;color:{mom_color}">{mom}</td>
                  <td style="padding:10px 14px;font-size:13px;color:rgba(255,255,255,.6)">{seg.get('yoy','')}</td>
                  <td style="padding:10px 14px;font-size:12px;color:rgba(255,255,255,.5)">{seg.get('active',''):,}</td>
                </tr>"""
        return f"""
        <div style="background:rgba(255,255,255,.07);border-radius:16px;overflow:auto;flex:1">
          <table style="width:100%;border-collapse:collapse">
            <thead><tr style="border-bottom:1px solid rgba(255,255,255,.12)">
              {''.join(f'<th style="padding:10px 14px;text-align:left;font-size:11px;font-weight:700;color:rgba(255,255,255,.4);letter-spacing:.06em;text-transform:uppercase">{h}</th>' for h in ['Market','Segment','RR','MoM','YoY','Active'])}
            </tr></thead>
            <tbody>{rows}</tbody>
          </table>
        </div>"""

    # Cancel rate segments
    if "cancel_rate_han" in content:
        cr = content["cancel_rate_han"]
        segs = cr.get("segments", [])
        rows = "".join(
            f'<tr><td style="padding:8px 14px;font-size:13px;font-weight:600;color:white">{s.get("segment","")}</td>'
            f'<td style="padding:8px 14px;font-size:13px;color:{"#EF4444" if float(s.get("cr","0%").rstrip("%"))>14 else "white"}">{s.get("cr","")}</td>'
            f'<td style="padding:8px 14px;font-size:12px;color:rgba(255,255,255,.5)">{s.get("cancel_per_active","")}</td></tr>'
            for s in segs
        )
        insight = cr.get("insight", "")
        return f"""
        <div style="background:rgba(255,255,255,.07);border-radius:16px;overflow:auto;flex:1">
          <div style="padding:14px 18px;font-size:12px;color:rgba(255,255,255,.5);border-bottom:1px solid rgba(255,255,255,.1)">{insight}</div>
          <table style="width:100%;border-collapse:collapse">
            <thead><tr style="border-bottom:1px solid rgba(255,255,255,.12)">
              {''.join(f'<th style="padding:8px 14px;text-align:left;font-size:11px;font-weight:700;color:rgba(255,255,255,.4);text-transform:uppercase">{h}</th>' for h in ['Segment','CR','Cancel/Active'])}
            </tr></thead>
            <tbody>{rows}</tbody>
          </table>
        </div>"""

    # Capacity segments
    if "capacity_han_m8" in content:
        cap = content["capacity_han_m8"]
        segs = cap.get("segments", [])
        rows = "".join(
            f'<tr><td style="padding:10px 14px;font-size:13px;font-weight:600;color:white">{s.get("segment","")}</td>'
            f'<td style="padding:10px 14px;font-size:13px;color:white">{s.get("actual",""):,}</td>'
            f'<td style="padding:10px 14px;font-size:13px;color:rgba(255,255,255,.6)">{s.get("target",""):,}</td>'
            f'<td style="padding:10px 14px;font-size:13px;font-weight:700;color:{"#EF4444" if float(s.get("pct_target","0%").rstrip("%"))<60 else "#10B981"}">{s.get("pct_target","")}</td>'
            f'<td style="padding:10px 14px"><span style="font-size:11px;padding:3px 10px;border-radius:10px;background:{"rgba(239,68,68,.2)" if "Nguy" in s.get("status","") else "rgba(251,191,36,.2)"};color:{"#FCA5A5" if "Nguy" in s.get("status","") else "#FCD34D"}">{s.get("status","")}</span></td>'
            f'</tr>'
            for s in segs
        )
        return f"""
        <div style="background:rgba(255,255,255,.07);border-radius:16px;overflow:auto;flex:1">
          <table style="width:100%;border-collapse:collapse">
            <thead><tr style="border-bottom:1px solid rgba(255,255,255,.12)">
              {''.join(f'<th style="padding:10px 14px;text-align:left;font-size:11px;font-weight:700;color:rgba(255,255,255,.4);text-transform:uppercase">{h}</th>' for h in ['Segment','Actual','Target','% Target','Status'])}
            </tr></thead>
            <tbody>{rows}</tbody>
          </table>
        </div>"""

    return ""


def _render_highlights_lowlights(slide: dict, img: str | None, b: dict) -> str:
    content = slide.get("content", {})
    highlights = content.get("highlights", [])
    lowlights = content.get("lowlights", [])
    actions = content.get("actions_executed", [])

    def items_html(items, color, icon):
        if not items:
            return ""
        rows = ""
        for it in items:
            # Handle nested {market, items} hoặc flat string
            if isinstance(it, dict):
                market = it.get("market", "")
                sub_items = it.get("items", [])
                if market:
                    rows += (f'<div style="margin-bottom:16px">'
                             f'<div style="font-size:11px;font-weight:700;color:{color};'
                             f'letter-spacing:.08em;text-transform:uppercase;margin-bottom:6px">'
                             f'{icon} {market}</div>')
                    for s in sub_items:
                        rows += (f'<div style="display:flex;gap:10px;margin-bottom:6px;padding-left:4px">'
                                 f'<span style="color:{color};flex-shrink:0">·</span>'
                                 f'<span style="font-size:14px;color:#1e293b;line-height:1.5">{s}</span>'
                                 f'</div>')
                    rows += '</div>'
                else:
                    text = " | ".join(str(v) for v in it.values() if v)
                    rows += (f'<div style="display:flex;gap:12px;margin-bottom:10px">'
                             f'<span style="color:{color};flex-shrink:0">{icon}</span>'
                             f'<span style="font-size:14px;color:#1e293b;line-height:1.5">{text}</span>'
                             f'</div>')
            else:
                rows += (f'<div style="display:flex;gap:12px;margin-bottom:10px">'
                         f'<span style="color:{color};flex-shrink:0">{icon}</span>'
                         f'<span style="font-size:14px;color:#1e293b;line-height:1.5">{it}</span>'
                         f'</div>')
        return rows

    hl_html = items_html(highlights, b["success"], "▲")
    ll_html = items_html(lowlights, b["danger"], "▼")

    # Actions executed box
    actions_html = ""
    if actions:
        action_rows = ""
        for a in actions:
            if isinstance(a, dict):
                act_text = a.get("action", "")
                result = a.get("result", "")
                status = a.get("status", "")
                status_color = b["success"] if "Hoàn" in status else b["accent"]
                action_rows += (
                    f'<div style="display:flex;gap:12px;margin-bottom:10px;align-items:flex-start">'
                    f'<span style="color:{b["primary"]};flex-shrink:0;font-weight:700">→</span>'
                    f'<div><div style="font-size:13px;color:#1e293b;font-weight:600">{act_text}</div>'
                    f'<div style="font-size:12px;color:{b["text_muted"]}">{result}'
                    f'{"  ·  " if result and status else ""}'
                    f'<span style="color:{status_color}">{status}</span></div></div></div>'
                )
        actions_html = f"""
        <div style="margin-top:20px;padding:16px 20px;background:rgba(14,65,116,.06);border-radius:14px;border-left:3px solid {b['primary']}">
          <div style="font-size:11px;font-weight:700;color:{b['primary']};letter-spacing:.08em;text-transform:uppercase;margin-bottom:10px">Actions Đã Thực Hiện</div>
          {action_rows}
        </div>"""

    return f"""
    <div style="width:100%;height:100%;background:{b['bg']};display:flex;flex-direction:column;padding:48px 64px">
      <div style="font-size:26px;font-weight:800;color:{b['primary']};margin-bottom:6px;line-height:1.3">{slide.get('slide_title','')}</div>
      <div style="font-size:13px;color:{b['text_muted']};margin-bottom:24px">{slide.get('key_message','')}</div>
      <div style="display:grid;grid-template-columns:1fr 1fr;gap:20px;flex:1;min-height:0">
        <div style="background:white;border-radius:20px;padding:28px 32px;border-top:4px solid {b['success']};overflow:auto">
          <div style="font-size:11px;font-weight:700;color:{b['success']};letter-spacing:.1em;text-transform:uppercase;margin-bottom:18px">Highlights</div>
          {hl_html}
        </div>
        <div style="background:white;border-radius:20px;padding:28px 32px;border-top:4px solid {b['danger']};overflow:auto">
          <div style="font-size:11px;font-weight:700;color:{b['danger']};letter-spacing:.1em;text-transform:uppercase;margin-bottom:18px">Lowlights</div>
          {ll_html}
        </div>
      </div>
      {actions_html}
    </div>"""


def _render_action_table(slide: dict, img: str | None, b: dict) -> str:
    content = slide.get("content", {})
    actions = content.get("actions", [])

    rows_html = ""
    for i, a in enumerate(actions):
        bg = "white" if i % 2 == 0 else "#F8FAFC"
        rows_html += f"""
        <tr style="background:{bg}">
          <td style="padding:16px 20px;font-size:14px;font-weight:600;color:{b['primary']}">{i+1}</td>
          <td style="padding:16px 20px;font-size:14px;color:{b['text_dark']}">{a.get('action','')}</td>
          <td style="padding:16px 20px;font-size:13px;color:{b['text_muted']}">{a.get('pic','')}</td>
          <td style="padding:16px 20px;font-size:13px;color:{b['accent']};font-weight:600">{a.get('deadline','')}</td>
          <td style="padding:16px 20px">
            <span style="background:{_status_color(a.get('status',''), b)};color:white;padding:4px 14px;border-radius:20px;font-size:12px;font-weight:600">
              {a.get('status','Open')}
            </span>
          </td>
        </tr>"""

    return f"""
    <div style="width:100%;height:100%;background:{b['bg']};display:flex;flex-direction:column;padding:56px 72px">
      <div style="font-size:28px;font-weight:800;color:{b['primary']};margin-bottom:8px">{slide.get('slide_title','')}</div>
      <div style="font-size:15px;color:{b['text_muted']};margin-bottom:32px">{slide.get('key_message','')}</div>
      <div style="background:white;border-radius:24px;overflow:hidden;flex:1;box-shadow:0 4px 24px rgba(0,0,0,.06)">
        <table style="width:100%;border-collapse:collapse">
          <thead>
            <tr style="background:{b['primary']}">
              <th style="padding:16px 20px;color:white;font-size:13px;text-align:left;width:40px">#</th>
              <th style="padding:16px 20px;color:white;font-size:13px;text-align:left">Action</th>
              <th style="padding:16px 20px;color:white;font-size:13px;text-align:left">PIC</th>
              <th style="padding:16px 20px;color:white;font-size:13px;text-align:left">Deadline</th>
              <th style="padding:16px 20px;color:white;font-size:13px;text-align:left">Status</th>
            </tr>
          </thead>
          <tbody>{rows_html}</tbody>
        </table>
      </div>
    </div>"""


def _render_executive_summary(slide: dict, img: str | None, b: dict) -> str:
    content = slide.get("content", {})
    metrics = content.get("metrics", [])
    findings = content.get("findings", [])

    metrics_html = ""
    for m in metrics[:4]:
        metrics_html += f"""
        <div style="text-align:center">
          <div style="font-size:48px;font-weight:900;color:{b['accent']};line-height:1;letter-spacing:-.02em">{m.get('value','')}</div>
          <div style="font-size:13px;color:rgba(255,255,255,.6);margin-top:6px">{m.get('label','')}</div>
        </div>"""

    findings_html = "".join(
        f'<div style="display:flex;gap:10px;margin-bottom:10px">'
        f'<span style="color:{b["accent"]};font-weight:700;flex-shrink:0">→</span>'
        f'<span style="font-size:14px;color:rgba(255,255,255,.85)">{f}</span>'
        f'</div>'
        for f in findings[:5]
    )

    img_section = ""
    if img and Path(img).exists():
        img_b64 = _img_to_base64(img)
        img_section = f'<img src="{img_b64}" style="position:absolute;right:0;top:0;height:100%;width:45%;object-fit:cover;opacity:.25">'

    return f"""
    <div style="width:100%;height:100%;background:linear-gradient(135deg,{b['primary']} 0%,#0a2a52 100%);
                display:flex;flex-direction:column;justify-content:center;padding:72px 80px;position:relative;overflow:hidden">
      {img_section}
      <div style="position:relative;z-index:1">
        <div style="font-size:13px;font-weight:600;color:{b['accent']};letter-spacing:.15em;text-transform:uppercase;margin-bottom:16px">
          Executive Summary
        </div>
        <div style="font-size:42px;font-weight:900;color:white;line-height:1.15;margin-bottom:40px;max-width:800px">
          {slide.get('slide_title','')}
        </div>
        <div style="display:flex;gap:48px;margin-bottom:48px;padding-bottom:40px;border-bottom:1px solid rgba(255,255,255,.15)">
          {metrics_html}
        </div>
        <div style="max-width:700px">{findings_html}</div>
      </div>
    </div>"""


def _render_problem_solution(slide: dict, img: str | None, b: dict) -> str:
    content = slide.get("content", {})

    def _text(v) -> str:
        if isinstance(v, str):
            return v
        if isinstance(v, list):
            return "<br>".join(f"· {i}" if isinstance(i, str) else f"· {' | '.join(str(x) for x in i.values())}" for i in v)
        if isinstance(v, dict):
            return "<br>".join(f"<b>{k}:</b> {_text(vv)}" for k, vv in v.items())
        return str(v)

    def box(label, val, color, icon):
        text = _text(val) if val else "—"
        return f"""
        <div style="background:white;border-radius:18px;padding:22px 24px;border-left:4px solid {color};overflow:auto">
          <div style="font-size:10px;font-weight:700;color:{color};letter-spacing:.1em;text-transform:uppercase;margin-bottom:10px">{icon} {label}</div>
          <div style="font-size:13px;color:{b['text_dark']};line-height:1.65">{text}</div>
        </div>"""

    # Nếu content có flat keys problem/solution dùng đó, nếu không flatten nested sections
    if content.get("problem") or content.get("solution"):
        boxes = [
            box("Problem", content.get("problem"), b["danger"], "⚠"),
            box("Root Cause", content.get("root_cause"), b["accent"], "🔍"),
            box("Solution", content.get("solution"), b["success"], "✓"),
            box("Impact", content.get("impact"), b["primary"], "→"),
        ]
    else:
        # Flatten nested sections (Bulky case)
        section_colors = [b["primary"], b["accent"], b["success"], b["danger"]]
        icons = ["📦", "🔧", "📋", "→"]
        boxes = [
            box(k.replace("_", " ").title(), v, section_colors[i % 4], icons[i % 4])
            for i, (k, v) in enumerate(content.items())
        ]

    grid = "\n".join(boxes[:4])
    return f"""
    <div style="width:100%;height:100%;background:{b['bg']};display:flex;flex-direction:column;padding:44px 56px">
      <div style="font-size:26px;font-weight:800;color:{b['primary']};margin-bottom:6px;line-height:1.3">{slide.get('slide_title','')}</div>
      <div style="font-size:13px;color:{b['text_muted']};margin-bottom:20px">{slide.get('key_message','')}</div>
      <div style="display:grid;grid-template-columns:1fr 1fr;grid-template-rows:1fr 1fr;gap:16px;flex:1;min-height:0">
        {grid}
      </div>
    </div>"""


def _render_default(slide: dict, img: str | None, b: dict) -> str:
    """Fallback layout cho các loại chưa có renderer riêng."""
    content = slide.get("content", {})
    points = content.get("points", content.get("items", []))

    points_html = "".join(
        f'<div style="display:flex;gap:14px;margin-bottom:16px;align-items:flex-start">'
        f'<div style="width:8px;height:8px;border-radius:50%;background:{b["accent"]};margin-top:7px;flex-shrink:0"></div>'
        f'<div style="font-size:16px;color:{b["text_dark"]};line-height:1.6">{p}</div>'
        f'</div>'
        for p in (points if isinstance(points, list) else [str(points)])
    )

    img_html = ""
    if img and Path(img).exists():
        img_b64 = _img_to_base64(img)
        img_html = f'<img src="{img_b64}" style="max-height:320px;max-width:100%;object-fit:contain;border-radius:16px;margin-top:24px">'

    return f"""
    <div style="width:100%;height:100%;background:{b['bg']};display:flex;flex-direction:column;padding:64px 80px">
      <div style="display:flex;align-items:center;gap:16px;margin-bottom:12px">
        <div style="width:6px;height:40px;background:{b['accent']};border-radius:4px"></div>
        <div style="font-size:30px;font-weight:800;color:{b['primary']}">{slide.get('slide_title','')}</div>
      </div>
      <div style="font-size:16px;color:{b['text_muted']};margin-bottom:40px;padding-left:22px">{slide.get('key_message','')}</div>
      <div style="flex:1;padding-left:22px">
        {points_html}
        {img_html}
      </div>
    </div>"""


def _render_section_divider(slide: dict, img: str | None, b: dict) -> str:
    content = slide.get("content", {})
    section_num = content.get("section_number", "")
    subtitle = content.get("subtitle", slide.get("key_message", ""))
    return f"""
    <div style="width:100%;height:100%;background:linear-gradient(135deg,{b['primary']} 0%,#1a3a6b 100%);
                display:flex;flex-direction:column;align-items:center;justify-content:center;text-align:center;padding:80px">
      {f'<div style="font-size:80px;font-weight:900;color:rgba(255,255,255,.1);line-height:1;margin-bottom:16px">{section_num}</div>' if section_num else ""}
      <div style="font-size:52px;font-weight:900;color:white;line-height:1.15;max-width:800px;margin-bottom:24px">
        {slide.get('slide_title','')}
      </div>
      {f'<div style="font-size:20px;color:rgba(255,255,255,.6);max-width:600px">{subtitle}</div>' if subtitle else ""}
      <div style="width:60px;height:4px;background:{b["accent"]};border-radius:2px;margin-top:40px"></div>
    </div>"""


LAYOUT_RENDERERS = {
    "kpi_cards": _render_kpi_cards,
    "highlights_lowlights": _render_highlights_lowlights,
    "action_table": _render_action_table,
    "executive_summary": _render_executive_summary,
    "problem_solution": _render_problem_solution,
    "section_divider": _render_section_divider,
    # Các layout khác fallback về _render_default
}


# ─── PPTX builder ─────────────────────────────────────────────────────────────

def _add_pptx_slide(slide, slide_data: dict, images_dir: str, b: dict, prs):
    """Thêm 1 slide vào PPTX với title + content text boxes."""
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    W = prs.slide_width
    H = prs.slide_height

    # Background
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(b["primary"].lstrip("#"))

    # Title box
    txBox = slide.shapes.add_textbox(Emu(int(W * 0.05)), Emu(int(H * 0.08)),
                                      Emu(int(W * 0.9)), Emu(int(H * 0.15)))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = slide_data.get("slide_title", "")
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # Key message
    km_box = slide.shapes.add_textbox(Emu(int(W * 0.05)), Emu(int(H * 0.25)),
                                       Emu(int(W * 0.9)), Emu(int(H * 0.08)))
    tf2 = km_box.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = slide_data.get("key_message", "")
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(200, 215, 235)

    # Content text
    content = slide_data.get("content", {})
    content_text = _content_to_text(content)
    if content_text:
        ct_box = slide.shapes.add_textbox(Emu(int(W * 0.05)), Emu(int(H * 0.35)),
                                           Emu(int(W * 0.9)), Emu(int(H * 0.55)))
        tf3 = ct_box.text_frame
        tf3.word_wrap = True
        for line in content_text.split("\n"):
            p3 = tf3.add_paragraph()
            p3.text = line
            p3.font.size = Pt(14)
            p3.font.color.rgb = RGBColor(230, 235, 245)

    # Image if available
    img_path = _get_image_path(slide_data, images_dir)
    if img_path and Path(img_path).exists():
        try:
            slide.shapes.add_picture(
                img_path,
                left=Emu(int(W * 0.55)), top=Emu(int(H * 0.3)),
                width=Emu(int(W * 0.4)), height=Emu(int(H * 0.55))
            )
        except Exception:
            pass  # ảnh lỗi thì bỏ qua


# ─── Utils ────────────────────────────────────────────────────────────────────

def _get_image_path(slide: dict, images_dir: str) -> str | None:
    num = slide.get("slide_num", 0)
    path = Path(images_dir) / f"slide_{num:02d}_image.png"
    return str(path) if path.exists() else None


def _img_to_base64(path: str) -> str:
    with open(path, "rb") as f:
        data = base64.b64encode(f.read()).decode()
    return f"data:image/png;base64,{data}"


def _status_color(status: str, b: dict) -> str:
    s = status.lower()
    if s in ("done", "completed", "closed"):
        return b["success"]
    if s in ("in progress", "ongoing"):
        return b["accent"]
    if s in ("blocked", "critical"):
        return b["danger"]
    return b["primary"]


def _content_to_text(content: dict) -> str:
    """Flatten content dict thành plain text cho PPTX."""
    lines = []
    for k, v in content.items():
        if isinstance(v, list):
            for item in v:
                if isinstance(item, dict):
                    lines.append("• " + " | ".join(str(x) for x in item.values()))
                else:
                    lines.append(f"• {item}")
        elif isinstance(v, str) and v:
            lines.append(f"{v}")
    return "\n".join(lines)
