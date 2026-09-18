#!/usr/bin/env python3
"""
Native PowerPoint PPTX Exporter Engine (Phase 10).
Converts SlideJSON into 16:9 native editable PowerPoint presentations (.pptx).
"""

import sys
import os
import json

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    Presentation = None


def export_slide_json_to_pptx(slide_json_path: str, output_pptx_path: str) -> str:
    with open(slide_json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    if Presentation is None:
        print("[WARNING] python-pptx library missing. Skipping native PPTX generation.")
        return output_pptx_path

    prs = Presentation()
    # Set 16:9 aspect ratio dimensions (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Theme colors
    PRIMARY_RGB = RGBColor(14, 65, 116)     # #0E4174
    ACCENT_RGB = RGBColor(255, 127, 50)    # #FF7F32
    TEXT_MAIN_RGB = RGBColor(15, 23, 42)    # #0F172A
    TEXT_MUTED_RGB = RGBColor(71, 85, 105)  # #475569
    CARD_BG_RGB = RGBColor(255, 255, 255)

    for slide_data in data.get("slides", []):
        slide = prs.slides.add_slide(blank_slide_layout)

        # 1. Header Takeaway Title Box
        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(12.133), Inches(1.2))
        tf = title_box.text_frame
        tf.word_wrap = True
        
        # Sub-header category tag
        p_tag = tf.paragraphs[0]
        p_tag.text = f"AHAMOVE DRIVER MANAGEMENT • SLIDE {slide_data.get('slide_id', 1)}"
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = ACCENT_RGB

        # Main takeaway header
        p_title = tf.add_paragraph()
        p_title.text = slide_data.get("takeaway_header", "Executive Summary")
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = PRIMARY_RGB

        # 2. Hero KPI Cards
        hero_kpis = slide_data.get("hero_kpis", [])
        if hero_kpis:
            kpi_count = len(hero_kpis)
            card_w = Inches(12.133 / kpi_count - 0.2)
            for idx, kpi in enumerate(hero_kpis):
                left_pos = Inches(0.6 + idx * (12.133 / kpi_count))
                shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(1.8), card_w, Inches(1.2))
                shape.fill.solid()
                shape.fill.fore_color.rgb = CARD_BG_RGB
                shape.line.color.rgb = PRIMARY_RGB

                ktf = shape.text_frame
                ktf.word_wrap = True
                kp1 = ktf.paragraphs[0]
                kp1.text = kpi.get("label", "").upper()
                kp1.font.size = Pt(10)
                kp1.font.color.rgb = TEXT_MUTED_RGB

                kp2 = ktf.add_paragraph()
                kp2.text = kpi.get("value", "")
                kp2.font.size = Pt(28)
                kp2.font.bold = True
                kp2.font.color.rgb = PRIMARY_RGB

        # 3. Content Glass Cards
        cards = slide_data.get("cards", [])
        top_offset = Inches(3.2) if hero_kpis else Inches(2.0)
        card_h = Inches(3.6) if hero_kpis else Inches(4.8)

        if cards:
            c_count = min(len(cards), 2)
            c_w = Inches(5.9)
            for c_idx in range(c_count):
                c_data = cards[c_idx]
                c_left = Inches(0.6 + c_idx * 6.233)
                
                c_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, top_offset, c_w, card_h)
                c_shape.fill.solid()
                c_shape.fill.fore_color.rgb = CARD_BG_RGB
                c_shape.line.color.rgb = RGBColor(226, 232, 240)

                ctf = c_shape.text_frame
                ctf.word_wrap = True
                
                # Card Title
                cp1 = ctf.paragraphs[0]
                cp1.text = f"{c_data.get('icon_badge', '💡')} {c_data.get('title', '')}"
                cp1.font.size = Pt(16)
                cp1.font.bold = True
                cp1.font.color.rgb = PRIMARY_RGB

                # Bullet points
                for bullet in c_data.get("bullets", []):
                    bp = ctf.add_paragraph()
                    bp.text = f"• {bullet}"
                    bp.font.size = Pt(12)
                    bp.font.color.rgb = TEXT_MAIN_RGB

        # 4. Footer Source Traceability
        footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(12.133), Inches(0.3))
        ftf = footer_box.text_frame
        fp = ftf.paragraphs[0]
        sources = slide_data.get("source_slides", [])
        fp.text = f"Confidential • Ahamove Internal Strategy | Source: Slide {', '.join(map(str, sources)) if sources else 'Data Engine'}"
        fp.font.size = Pt(9)
        fp.font.color.rgb = TEXT_MUTED_RGB

    os.makedirs(os.path.dirname(output_pptx_path), exist_ok=True)
    prs.save(output_pptx_path)
    return output_pptx_path


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python3 export_pptx.py <slide_json_path> <output_pptx_path>")
        sys.exit(1)
    
    json_path = sys.argv[1]
    output_pptx = sys.argv[2]
    res = export_slide_json_to_pptx(json_path, output_pptx)
    print(f"Exported Native PPTX 16:9 Presentation: {res}")
