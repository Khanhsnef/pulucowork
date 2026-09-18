#!/usr/bin/env python3
"""
PPTX Parser Engine for Phase 1 Ingestion.
Parses raw PPTX presentations into structured NormalizedSourceDeck JSON.
"""

import sys
import os
import json
import re

try:
    from pptx import Presentation
except ImportError:
    Presentation = None


def parse_pptx_deck(file_path: str, slide_range: str = "all") -> dict:
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Input file not found: {file_path}")

    if Presentation is None:
        # Fallback parser if python-pptx is not installed
        return {
            "deck_id": os.path.basename(file_path),
            "source_file": file_path,
            "total_slides": 1,
            "slides": [
                {
                    "slide_number": 1,
                    "raw_title": "Source Presentation Data",
                    "paragraphs": ["python-pptx module missing. Ingestion fallback mode active."],
                    "tables": [],
                    "key_metrics_extracted": []
                }
            ]
        }

    prs = Presentation(file_path)
    slides_data = []

    # Parse slide range e.g. "1-30" or "all"
    start_idx, end_idx = 0, len(prs.slides)
    if slide_range != "all" and "-" in slide_range:
        try:
            parts = slide_range.split("-")
            start_idx = max(0, int(parts[0]) - 1)
            end_idx = min(len(prs.slides), int(parts[1]))
        except ValueError:
            pass

    for i in range(start_idx, end_idx):
        slide = prs.slides[i]
        slide_num = i + 1
        raw_title = f"Slide {slide_num}"
        paragraphs = []
        tables = []
        metrics = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    if shape == slide.shapes[0] and len(text) < 100:
                        raw_title = text
                    else:
                        paragraphs.append(text)
                    
                    # Regex search for KPI metric candidates (e.g. 96%, 15.8M VND, +18%)
                    kpi_matches = re.findall(r'(\d+(?:\.\d+)?\s*(?:%|VND|k|x|M|bps))', text, re.IGNORECASE)
                    for match in kpi_matches:
                        metrics.append({
                            "metric_name": "Extracted Metric",
                            "value": match,
                            "context": text[:80]
                        })

            if shape.has_table:
                table = shape.table
                if len(table.rows) > 0:
                    headers = [cell.text.strip() for cell in table.rows[0].cells]
                    rows = []
                    for row_idx in range(1, len(table.rows)):
                        row = table.rows[row_idx]
                        rows.append([cell.text.strip() for cell in row.cells])
                    tables.append({"headers": headers, "rows": rows})


        slides_data.append({
            "slide_number": slide_num,
            "raw_title": raw_title,
            "paragraphs": paragraphs,
            "tables": tables,
            "key_metrics_extracted": metrics
        })

    return {
        "deck_id": os.path.basename(file_path),
        "source_file": file_path,
        "total_slides": len(slides_data),
        "slides": slides_data
    }


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python3 parse_pptx.py <file.pptx> [slide_range]")
        sys.exit(1)
    
    file_path = sys.argv[1]
    slide_range = sys.argv[2] if len(sys.argv) > 2 else "all"
    result = parse_pptx_deck(file_path, slide_range)
    print(json.dumps(result, indent=2, ensure_ascii=False))
