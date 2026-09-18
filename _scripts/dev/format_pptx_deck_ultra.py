import sys
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

# Theme Palette: Ultra-Sleek Glossy Executive Dark Theme
BG_COLOR = RGBColor(11, 19, 44)           # #0B132C Deep Executive Dark Canvas
CARD_BG = RGBColor(24, 36, 66)            # #182442 Glass Dark Slate Card
CARD_BG_ALT = RGBColor(16, 25, 48)       # #101930 Darker Sub-card
CARD_BORDER = RGBColor(50, 75, 120)       # #324B78 Metallic Border Stroke

TEXT_WHITE = RGBColor(255, 255, 255)
TEXT_BRIGHT = RGBColor(241, 245, 249)     # #F1F5F9 Crisp Off-white
TEXT_MUTED = RGBColor(160, 174, 192)      # #A0AEC0 Bright Muted Silver

ORANGE_BRAND = RGBColor(255, 115, 0)      # #FF7300 Vivid Ahamove Glossy Orange
ORANGE_CARD = RGBColor(40, 25, 15)
BLUE_HEADER = RGBColor(26, 54, 138)       # #1A368A Royal Glossy Dark Blue
BLUE_ACTIVE = RGBColor(37, 99, 235)       # #2563EB Vibrant Blue Pill
CYAN_ACCENT = RGBColor(56, 189, 248)      # #38BDF8 Glowing Cyan

RED_ACCENT = RGBColor(239, 68, 68)        # #EF4444 Glossy Red
RED_BG = RGBColor(55, 25, 35)

AMBER_ACCENT = RGBColor(245, 158, 11)     # #F59E0B Glossy Amber
AMBER_BG = RGBColor(50, 38, 20)

GREEN_ACCENT = RGBColor(16, 185, 129)     # #10B981 Glossy Emerald
GREEN_BG = RGBColor(20, 50, 40)

prs = pptx.Presentation('/Users/ts-1148/Desktop/Pulu-workspace/Data sources/BaoCao_ThangDM_T7_Core_v7_1.pptx')

def style_text_runs(shape, font_family="Arial", default_color=TEXT_BRIGHT):
    if not shape.has_text_frame:
        return
    for p in shape.text_frame.paragraphs:
        p.font.name = font_family
        for r in p.runs:
            r.font.name = font_family
            # If text color is dark or missing, set to default bright text
            if r.font.color and r.font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB:
                rgb = r.font.color.rgb
                if rgb[0] < 120 and rgb[1] < 120 and rgb[2] < 120:
                    r.font.color.rgb = default_color
            else:
                r.font.color.rgb = default_color

def style_card_shape(shape, fill_rgb, border_rgb=CARD_BORDER, line_width=1):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    if border_rgb:
        shape.line.color.rgb = border_rgb
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()

for slide_idx, slide in enumerate(prs.slides):
    # Set slide dark canvas background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR
    
    for shape in slide.shapes:
        # 1. Handle Card Shapes & Fills
        if shape.shape_type in [MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM]:
            if shape.fill.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB:
                rgb = shape.fill.fore_color.rgb
                # Light grey / white container cards (#EEEEEE or white) -> Sleek Glass Dark Card
                if rgb[0] > 200 and rgb[1] > 200 and rgb[2] > 200:
                    style_card_shape(shape, CARD_BG, CARD_BORDER, 1)
                # Dark navy header / accent cards
                elif rgb[0] < 45 and rgb[1] < 65 and rgb[2] < 100:
                    # Check if it's the active tab in project slides (Slide 4-7)
                    if slide_idx in [3, 4, 5, 6] and shape.top and shape.top.inches < 1.0:
                        style_card_shape(shape, ORANGE_BRAND, ORANGE_BRAND, 1)
                    else:
                        style_card_shape(shape, BLUE_HEADER, ORANGE_BRAND, 1.5)
                # Orange header lines / pills
                elif rgb[0] > 220 and rgb[1] > 90 and rgb[2] < 50:
                    style_card_shape(shape, ORANGE_BRAND, None)
                # Red accent pills / cards
                elif rgb[0] > 200 and rgb[1] < 80 and rgb[2] < 80:
                    style_card_shape(shape, RED_ACCENT, None)
                # Green accent pills / cards
                elif rgb[0] < 50 and rgb[1] > 150 and rgb[2] < 100:
                    style_card_shape(shape, GREEN_ACCENT, None)
                # Amber accent pills / cards
                elif rgb[0] > 200 and rgb[1] > 140 and rgb[2] < 50:
                    style_card_shape(shape, AMBER_ACCENT, None)
                    
        # 2. Handle Text Framing & Typography
        if shape.has_text_frame:
            style_text_runs(shape, font_family="Arial", default_color=TEXT_BRIGHT)
            
            txt = shape.text_frame.text.strip()
            
            # Slide Title Header
            if shape.top and shape.top.inches < 1.0 and shape.height and shape.height.inches > 0.4:
                for p in shape.text_frame.paragraphs:
                    p.font.bold = True
                    for r in p.runs:
                        r.font.color.rgb = TEXT_WHITE
                        
            # Kicker tag (top-left)
            if shape.top and shape.top.inches < 0.35 and shape.left and shape.left.inches < 1.0:
                for p in shape.text_frame.paragraphs:
                    p.font.bold = True
                    for r in p.runs:
                        r.font.color.rgb = ORANGE_BRAND
                        
            # Subtitle / Presenter info (top-right)
            if shape.top and shape.top.inches < 0.35 and shape.left and shape.left.inches > 7.0:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = TEXT_MUTED
                        
            # Active Project Tab Text (Slide 4-7)
            if slide_idx in [3, 4, 5, 6] and shape.top and shape.top.inches < 1.0 and shape.width and shape.width.inches < 3.0:
                for p in shape.text_frame.paragraphs:
                    p.font.bold = True
                    for r in p.runs:
                        if r.font.color and r.font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB:
                            if r.font.color.rgb == ORANGE_BRAND:
                                r.font.color.rgb = TEXT_WHITE
                                
            # Key Takeaway keyword accent
            if "Key Takeaway" in txt:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if "Key Takeaway" in r.text:
                            r.font.color.rgb = ORANGE_BRAND
                            r.font.bold = True
                            
            # Numbers / Stat Callouts in Card 1 (Slides 4-7)
            if slide_idx in [3, 4, 5, 6] and txt in ["~2x", "68.2%", "-10%", "2/9"]:
                for p in shape.text_frame.paragraphs:
                    p.font.size = Pt(38)
                    p.font.bold = True
                    for r in p.runs:
                        r.font.color.rgb = TEXT_WHITE
                        
        # 3. Tables (Slide 3 & Slide 9)
        if shape.has_table:
            table = shape.table
            for row_idx, row in enumerate(table.rows):
                for cell_idx, cell in enumerate(row.cells):
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    # Header row
                    if row_idx == 0:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = BLUE_HEADER
                        for p in cell.text_frame.paragraphs:
                            p.alignment = PP_ALIGN.LEFT
                            p.font.name = "Arial"
                            p.font.bold = True
                            for r in p.runs:
                                r.font.color.rgb = TEXT_WHITE
                    else:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = CARD_BG if row_idx % 2 == 1 else CARD_BG_ALT
                        for p in cell.text_frame.paragraphs:
                            p.alignment = PP_ALIGN.LEFT
                            p.font.name = "Arial"
                            for r in p.runs:
                                txt_c = r.text.strip()
                                if "ĐỎ" in txt_c:
                                    r.font.color.rgb = RED_ACCENT
                                    r.font.bold = True
                                elif "VÀNG" in txt_c:
                                    r.font.color.rgb = AMBER_ACCENT
                                    r.font.bold = True
                                elif "XANH" in txt_c:
                                    r.font.color.rgb = GREEN_ACCENT
                                    r.font.bold = True
                                else:
                                    # Ensure high contrast text inside dark table cells
                                    r.font.color.rgb = TEXT_BRIGHT

output_file = '/Users/ts-1148/Desktop/Pulu-workspace/Data sources/BaoCao_ThangDM_T7_Core_v7_1.pptx'
prs.save(output_file)
print("Ultra formatted presentation saved to:", output_file)
