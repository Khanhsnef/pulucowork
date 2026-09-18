import sys
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

# Theme Palette: Executive Light Mode (Refined)
BG_LIGHT = RGBColor(248, 250, 252)         # #F8FAFC Off-White Canvas
CARD_WHITE = RGBColor(255, 255, 255)       # Pure White Card Fill
CARD_BORDER = RGBColor(203, 213, 225)      # #CBD5E1 Clean Border Stroke
CARD_BORDER_SOFT = RGBColor(226, 232, 240) # #E2E8F0 Soft Border

TEXT_DARK = RGBColor(15, 23, 42)           # #0F172A Deep Slate Heading Text
TEXT_BODY = RGBColor(51, 65, 85)           # #334155 High-readability Charcoal Body
TEXT_MUTED = RGBColor(100, 116, 139)       # #64748B Subdued Silver/Slate
TEXT_WHITE = RGBColor(255, 255, 255)

ORANGE_BRAND = RGBColor(255, 102, 0)       # #FF6600 Ahamove Vivid Orange
NAVY_HEADER = RGBColor(15, 23, 42)         # #0F172A Deep Navy Header
TAB_INACTIVE = RGBColor(241, 245, 249)     # #F1F5F9 Light Ice Slate Inactive Tab

RED_ACCENT = RGBColor(220, 38, 38)         # #DC2626 Red Accent Text
AMBER_ACCENT = RGBColor(217, 119, 6)       # #D97706 Amber Accent Text
GREEN_ACCENT = RGBColor(5, 150, 105)       # #059669 Green Accent Text

FONT_NAME = "Helvetica Neue"

prs = pptx.Presentation('/Users/ts-1148/Desktop/Pulu-workspace/Data sources/BaoCao_ThangDM_T7_Core_v7_1.pptx')

def make_subtle_rounded(shape, corner_val=2500):
    if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        spPr = shape.element.spPr
        prstGeom = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom')
        if prstGeom is not None:
            prstGeom.set('prst', 'roundRect')
            avLst = prstGeom.find('{http://schemas.openxmlformats.org/drawingml/2006/main}avLst')
            if avLst is None:
                avLst = OxmlElement('a:avLst')
                prstGeom.append(avLst)
            else:
                avLst.clear()
            gd = parse_xml(f'<a:gd {nsdecls("a")} name="adj" fmla="val {corner_val}"/>')
            avLst.append(gd)

def hide_shape_border_and_fill(shape):
    shape.fill.background()
    shape.line.fill.background()

def style_card_container(shape, fill_rgb=CARD_WHITE, border_rgb=CARD_BORDER, line_width=1, corner_val=2500):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    if border_rgb:
        shape.line.color.rgb = border_rgb
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    make_subtle_rounded(shape, corner_val)

# Tab active left boundaries for Slides 4, 5, 6, 7 (0-indexed slides 3, 4, 5, 6)
SLIDE_ACTIVE_TAB_LEFT = {
    3: (0.0, 2.0),     # Slide 4: MiniHub (left ~ 0.60 in)
    4: (3.0, 5.0),     # Slide 5: EV (left ~ 3.65 in)
    5: (6.0, 8.0),     # Slide 6: Baga Bulky (left ~ 6.71 in)
    6: (9.0, 11.5)     # Slide 7: Driver Journey (left ~ 9.76 in)
}

for slide_idx, slide in enumerate(prs.slides):
    # Set light canvas background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG_LIGHT
    
    for shape in slide.shapes:
        top_in = shape.top / 914400.0 if shape.top else 0
        left_in = shape.left / 914400.0 if shape.left else 0
        w_in = shape.width / 914400.0 if shape.width else 0
        h_in = shape.height / 914400.0 if shape.height else 0

        # Apply font family across all text frames
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                p.font.name = FONT_NAME
                for r in p.runs:
                    r.font.name = FONT_NAME

        # -------------------------------------------------------------
        # 1. TOP PROGRESS NAVIGATION TABS (Slides 4-7)
        # -------------------------------------------------------------
        if slide_idx in [3, 4, 5, 6] and top_in < 0.85 and w_in > 2.0 and w_in < 3.5:
            active_range = SLIDE_ACTIVE_TAB_LEFT[slide_idx]
            is_active = (active_range[0] <= left_in <= active_range[1])
            
            # If shape is thin accent line (h < 0.1"), hide it so there are no double lines
            if h_in < 0.1:
                hide_shape_border_and_fill(shape)
                continue
                
            if is_active:
                style_card_container(shape, ORANGE_BRAND, ORANGE_BRAND, 1, corner_val=2000)
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        p.font.bold = True
                        for r in p.runs:
                            r.font.color.rgb = TEXT_WHITE
            else:
                style_card_container(shape, TAB_INACTIVE, CARD_BORDER, 1, corner_val=2000)
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        p.font.bold = False
                        for r in p.runs:
                            r.font.color.rgb = TEXT_MUTED
            continue

        # -------------------------------------------------------------
        # 2. ACCENT LINES ABOVE CARDS (Hide thin duplicate line shapes)
        # -------------------------------------------------------------
        if h_in < 0.1 and top_in > 1.5 and top_in < 7.0 and w_in > 1.0:
            hide_shape_border_and_fill(shape)
            continue

        # -------------------------------------------------------------
        # 3. CARD CONTAINERS & SHAPES
        # -------------------------------------------------------------
        if shape.shape_type in [MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM]:
            # Convert rectangles to subtle rounded rectangles
            spPr = shape.element.spPr
            prstGeom = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom')
            if prstGeom is not None and prstGeom.get('prst') == 'rect':
                make_subtle_rounded(shape, corner_val=2500)

            if shape.fill.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB:
                rgb = shape.fill.fore_color.rgb
                
                # Big white background card containers
                if (rgb[0] > 180 and rgb[1] > 180 and rgb[2] > 180) or (rgb[0] < 50 and rgb[1] < 60 and rgb[2] < 140):
                    style_card_container(shape, CARD_WHITE, CARD_BORDER, 1, corner_val=2500)
                # Orange accent pill / icons
                elif rgb[0] > 220 and rgb[1] > 90 and rgb[2] < 50:
                    style_card_container(shape, ORANGE_BRAND, None, 0, corner_val=2000)
                # Red accent pill / icons
                elif rgb[0] > 200 and rgb[1] < 80 and rgb[2] < 80:
                    style_card_container(shape, RED_ACCENT, None, 0, corner_val=2000)
                # Green accent pill / icons
                elif rgb[0] < 50 and rgb[1] > 150 and rgb[2] < 100:
                    style_card_container(shape, GREEN_ACCENT, None, 0, corner_val=2000)

        # -------------------------------------------------------------
        # 4. TYPOGRAPHY & TEXT ALIGNMENTS
        # -------------------------------------------------------------
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            
            # Default text color mapping for light backgrounds
            for p in shape.text_frame.paragraphs:
                p.font.name = FONT_NAME
                for r in p.runs:
                    r.font.name = FONT_NAME
                    if r.font.color and r.font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB:
                        c_rgb = r.font.color.rgb
                        if c_rgb[0] > 180 and c_rgb[1] > 180 and c_rgb[2] > 180:
                            r.font.color.rgb = TEXT_BODY
                        elif c_rgb[0] < 80 and c_rgb[1] < 80 and c_rgb[2] < 80:
                            r.font.color.rgb = TEXT_BODY
                    else:
                        r.font.color.rgb = TEXT_BODY

            # Slide Main Title Header
            if top_in < 1.0 and h_in > 0.4:
                for p in shape.text_frame.paragraphs:
                    p.font.name = FONT_NAME
                    p.font.bold = True
                    for r in p.runs:
                        r.font.name = FONT_NAME
                        r.font.color.rgb = TEXT_DARK
                        
            # Category Kicker (top-left)
            if top_in < 0.35 and left_in < 1.0:
                for p in shape.text_frame.paragraphs:
                    p.font.name = FONT_NAME
                    p.font.bold = True
                    for r in p.runs:
                        r.font.name = FONT_NAME
                        r.font.color.rgb = ORANGE_BRAND
                        
            # Subtitle / Presenter info (top-right)
            if top_in < 0.35 and left_in > 7.0:
                for p in shape.text_frame.paragraphs:
                    p.font.name = FONT_NAME
                    for r in p.runs:
                        r.font.name = FONT_NAME
                        r.font.color.rgb = TEXT_MUTED

            # Key Takeaway keyword accent
            if "Key Takeaway" in txt:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.name = FONT_NAME
                        if "Key Takeaway" in r.text:
                            r.font.color.rgb = ORANGE_BRAND
                            r.font.bold = True
                        else:
                            r.font.color.rgb = TEXT_BODY

            # Highlights / Lowlights card section titles
            if txt in ["HIGHLIGHTS", "LOWLIGHTS", "ĐÃ LÀM GÌ (T7)", "TÁC ĐỘNG", "BƯỚC TIẾP THEO", "CẦN SUPPORT TỪ"]:
                for p in shape.text_frame.paragraphs:
                    p.font.name = FONT_NAME
                    p.font.bold = True
                    for r in p.runs:
                        r.font.name = FONT_NAME
                        if txt == "HIGHLIGHTS":
                            r.font.color.rgb = GREEN_ACCENT
                        elif txt in ["LOWLIGHTS", "TÁC ĐỘNG"]:
                            r.font.color.rgb = RED_ACCENT
                        else:
                            r.font.color.rgb = TEXT_DARK

            # Large Stat Callouts in Card 1 (Slides 4-7)
            if slide_idx in [3, 4, 5, 6] and txt in ["~2x", "68.2%", "-10%", "2/9", "2 / 9"]:
                for p in shape.text_frame.paragraphs:
                    p.font.name = FONT_NAME
                    p.font.size = Pt(36)
                    p.font.bold = True
                    for r in p.runs:
                        r.font.name = FONT_NAME
                        r.font.color.rgb = TEXT_DARK

        # -------------------------------------------------------------
        # 5. TABLES (Slide 3 & Slide 9)
        # -------------------------------------------------------------
        if shape.has_table:
            table = shape.table
            for row_idx, row in enumerate(table.rows):
                for cell_idx, cell in enumerate(row.cells):
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    if row_idx == 0:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = NAVY_HEADER
                        for p in cell.text_frame.paragraphs:
                            p.alignment = PP_ALIGN.LEFT
                            p.font.name = FONT_NAME
                            p.font.bold = True
                            for r in p.runs:
                                r.font.name = FONT_NAME
                                r.font.color.rgb = TEXT_WHITE
                    else:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = CARD_WHITE if row_idx % 2 == 1 else BG_LIGHT
                        for p in cell.text_frame.paragraphs:
                            p.alignment = PP_ALIGN.LEFT
                            p.font.name = FONT_NAME
                            for r in p.runs:
                                r.font.name = FONT_NAME
                                txt_c = r.text.strip()
                                if "ĐỎ" in txt_c:
                                    r.font.color.rgb = RED_ACCENT
                                    r.font.bold = True
                                elif "VÀNG" in txt_c:
                                    r.font.color.rgb = AMBER_ACCENT
                                    r.font.bold = True
                                elif "XANH" in txt_c:
                                    r.font.color.rgb = GREEN_ACCENT
                                    r.font.bold = True
                                else:
                                    r.font.color.rgb = TEXT_BODY

output_file = '/Users/ts-1148/Desktop/Pulu-workspace/Data sources/BaoCao_ThangDM_T7_Core_v7_1.pptx'
prs.save(output_file)
print("Perfect light theme presentation saved successfully:", output_file)
