import asyncio
import os
import sys
from pptx import Presentation
from pptx.util import Inches
from playwright.async_api import async_playwright

async def render_html_slides_to_pptx(html_path, output_pptx_path):
    if not os.path.exists(html_path):
        print(f"Error: {html_path} does not exist")
        return

    temp_dir = os.path.join(os.path.dirname(html_path), "temp_slide_img")
    os.makedirs(temp_dir, exist_ok=True)

    file_url = f"file://{os.path.abspath(html_path)}"

    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={"width": 1280, "height": 720})
        await page.goto(file_url, wait_until="networkidle")

        # Get total slides count
        slides_count = await page.evaluate("document.querySelectorAll('.slide').length")
        print(f"Detected {slides_count} slides in HTML.")

        img_paths = []
        for i in range(slides_count):
            # Switch to slide i
            await page.evaluate(f"showSlide({i})")
            # Wait for Chart.js animation
            await page.wait_for_timeout(400)

            img_path = os.path.join(temp_dir, f"slide_{i+1}.png")
            # Screenshot stage element specifically
            stage_elem = await page.query_selector("#stage")
            if stage_elem:
                await stage_elem.screenshot(path=img_path, type="png")
            else:
                await page.screenshot(path=img_path, type="png")

            img_paths.append(img_path)
            print(f"Captured slide {i+1}/{slides_count} -> {img_path}")

        await browser.close()

    # Build 16:9 Presentation with rendered images
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for img_p in img_paths:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(img_p, 0, 0, width=Inches(13.333), height=Inches(7.5))

    prs.save(output_pptx_path)
    print(f"✅ Created Pixel-Perfect 16:9 PPTX at: {output_pptx_path}")

if __name__ == "__main__":
    html_file = "/Users/ts-1148/Desktop/Pulu-workspace/Output/Ahamove/04. OPS_METRICS/ahamove_master_ops_performance_report-169-slides.html"
    output_pptx = "/Users/ts-1148/Desktop/Pulu-workspace/Output/Ahamove/04. OPS_METRICS/ahamove_master_ops_performance_report-pixel-perfect.pptx"
    asyncio.run(render_html_slides_to_pptx(html_file, output_pptx))
