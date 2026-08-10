import sys
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

# Theme Palette: Ultra-Sleek Glossy Executive Dark Theme
BG_COLOR = RGBColor(11, 19, 44)           # #0B132C Deep Executive Dark
CARD_BG = RGBColor(22, 33, 62)            # #16213E Glass Dark Slate
CARD_BG_ALT = RGBColor(15, 23, 42)       # #0F172A Dark Sub-card
CARD_BORDER = RGBColor(40, 60, 95)        # #283C5F Subtle Metallic Line

TEXT_WHITE = RGBColor(255, 255, 255)
TEXT_MAIN = RGBColor(226, 232, 240)      # #E2E8F0 High contrast light
TEXT_MUTED = RGBColor(148, 163, 184)     # #94A3B8 Silver muted

ORANGE_BRAND = RGBColor(255, 115, 0)      # #FF7300 Vivid Ahamove Glossy Orange
ORANGE_LIGHT = RGBColor(255, 160, 60)
BLUE_HEADER = RGBColor(30, 58, 138)       # #1E3A8A Royal Dark Blue
CYAN_ACCENT = RGBColor(56, 189, 248)      # #38BDF8 Glowing Cyan

RED_ACCENT = RGBColor(239, 68, 68)        # #EF4444 Glossy Red
AMBER_ACCENT = RGBColor(245, 158, 11)     # #F59E0B Glossy Amber
GREEN_ACCENT = RGBColor(16, 185, 129)     # #10B981 Glossy Emerald

prs = pptx.Presentation('/Users/ts-1148/Desktop/Pulu-workspace/Data sources/BaoCao_ThangDM_T7_Core_v7_1.pptx')

def apply_text_styling(shape, font_family="Arial", default_color=TEXT_MAIN):
    if not shape.has_text_frame:
        return
    for p in shape.text_frame.paragraphs:
        if font_family:
            p.font.name = font_family
        for run in p.runs:
            if font_family:
                run.font.name = font_family
            # If text color is pure black or dark grey, convert to high-contrast TEXT_MAIN
            if run.font.color and run.font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB:
                rgb = run.font.color.rgb
                if rgb[0] < 80 and rgb[1] < 80 and rgb[2] < 80:
                    run.font.color.rgb = default_color
            elif not run.font.color or run.font.color.type != pptx.enum.dml.MSO_COLOR_TYPE.RGB:
                run.font.color.rgb = default_color

def style_card_shape(shape, fill_rgb, border_rgb=CARD_BORDER, line_width=1):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    if border_rgb:
        shape.line.color.rgb = border_rgb
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()

for slide_idx, slide in enumerate(prs.slides):
    # Set slide dark background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR
    
    for s_idx, shape in enumerate(slide.shapes):
        # 1. Handle background card shapes (rectangles, rounded rects)
        if shape.shape_type in [MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM]:
            if shape.fill.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB:
                rgb = shape.fill.fore_color.rgb
                # Light grey/white cards (#EEEEEE or white) -> Sleek dark card
                if rgb[0] > 200 and rgb[1] > 200 and rgb[2] > 200:
                    style_card_shape(shape, CARD_BG, CARD_BORDER, 1)
                # Dark navy accent containers -> Royal Blue with Orange border
                elif rgb[0] < 40 and rgb[1] < 60 and rgb[2] < 90:
                    style_card_shape(shape, BLUE_HEADER, ORANGE_BRAND, 1.5)
                # Orange accent lines or headers
                elif rgb[0] > 220 and rgb[1] > 90 and rgb[2] < 50:
                    style_card_shape(shape, ORANGE_BRAND, None)
                # Red accent headers/lines
                elif rgb[0] > 200 and rgb[1] < 80 and rgb[2] < 80:
                    style_card_shape(shape, RED_ACCENT, None)
                # Green accent headers/lines
                elif rgb[0] < 50 and rgb[1] > 150 and rgb[2] < 100:
                    style_card_shape(shape, GREEN_ACCENT, None)
                # Amber accent headers/lines
                elif rgb[0] > 200 and rgb[1] > 140 and rgb[2] < 50:
                    style_card_shape(shape, AMBER_ACCENT, None)
                    
        # 2. Handle Text Styling
        if shape.has_text_frame:
            apply_text_styling(shape, font_family="Arial", default_color=TEXT_MAIN)
            
            # Specific heading recoloring based on text content
            txt = shape.text_frame.text.strip()
            
            # Header Title (top title of slide)
            if shape.top and shape.top.inches < 1.0 and shape.height and shape.height.inches > 0.4:
                for p in shape.text_frame.paragraphs:
                    p.font.name = "Arial"
                    p.font.bold = True
                    for r in p.runs:
                        r.font.color.rgb = TEXT_WHITE
                        
            # Kicker tag (top left small text like "SO SÁNH CHẾ ĐỘ · 2 DỰ ÁN")
            if shape.top and shape.top.inches < 0.35 and shape.left and shape.left.inches < 1.0:
                for p in shape.text_frame.paragraphs:
                    p.font.name = "Arial"
                    p.font.bold = True
                    for r in p.runs:
                        r.font.color.rgb = ORANGE_BRAND
                        
            # Subtitle / Presenter (top right small text)
            if shape.top and shape.top.inches < 0.35 and shape.left and shape.left.inches > 7.0:
                for p in shape.text_frame.paragraphs:
                    p.font.name = "Arial"
                    for r in p.runs:
                        r.font.color.rgb = TEXT_MUTED
                        
            # Highlight orange keywords if needed
            if "Key Takeaway" in txt:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if "Key Takeaway" in r.text:
                            r.font.color.rgb = ORANGE_BRAND
                            r.font.bold = True
                            
        # 3. Handle Table Styling (Slide 3 and Slide 9)
        if shape.has_table:
            table = shape.table
            for row_idx, row in enumerate(table.rows):
                for cell_idx, cell in enumerate(row.cells):
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    # Header row
                    if row_idx == 0:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = BLUE_HEADER
                        for p in cell.text_frame.paragraphs:
                            p.alignment = PP_ALIGN.LEFT
                            p.font.name = "Arial"
                            p.font.bold = True
                            for r in p.runs:
                                r.font.color.rgb = TEXT_WHITE
                    else:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = CARD_BG if row_idx % 2 == 1 else CARD_BG_ALT
                        for p in cell.text_frame.paragraphs:
                            p.font.name = "Arial"
                            for r in p.runs:
                                txt_c = r.text.strip()
                                if "ĐỎ" in txt_c:
                                    r.font.color.rgb = RED_ACCENT
                                    r.font.bold = True
                                elif "VÀNG" in txt_c:
                                    r.font.color.rgb = AMBER_ACCENT
                                    r.font.bold = True
                                elif "XANH" in txt_c:
                                    r.font.color.rgb = GREEN_ACCENT
                                    r.font.bold = True
                                else:
                                    if not r.font.color or r.font.color.type != pptx.enum.dml.MSO_COLOR_TYPE.RGB or r.font.color.rgb[0] < 80:
                                        r.font.color.rgb = TEXT_MAIN

output_file = '/Users/ts-1148/Desktop/Pulu-workspace/BaoCao_ThangDM_T7_Core_v7_1_Formatted.pptx'
prs.save(output_file)
print("Updated presentation saved successfully:", output_file)
